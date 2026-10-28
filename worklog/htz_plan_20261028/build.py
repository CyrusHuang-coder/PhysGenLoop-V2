from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
import hashlib,json
from lxml import etree as E
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]; OUT=ROOT/'docs'; QA=Path(__file__).parent
sources=[OUT/'ESRT_立项申请表_huangtianzhi.docx',OUT/'指代式图像编辑空间控制决策.docx']
hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}
(QA/'source_hashes.json').write_text(json.dumps(hashes,ensure_ascii=False,indent=2),encoding='utf-8')

def setup(title):
 d=Document();s=d.sections[0];s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(2.3);s.right_margin=Cm(2.3)
 for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 s.page_width=Cm(21);s.page_height=Cm(29.7)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  st=d.styles[name];st.font.name='Microsoft YaHei';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=1.25;d.styles['Normal'].paragraph_format.space_after=Pt(6)
 d.styles['Title'].font.size=Pt(21);d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 2'].font.size=Pt(12)
 f=s.footer.paragraphs[0];f.alignment=2;f.add_run('导师讨论稿  ·  2026年10月28日  ·  ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 d.add_paragraph(title,'Title');return d
def p(d,text):d.add_paragraph(text)
def h(d,text):d.add_paragraph(text,'Heading 2')
def page(d,title):d.add_page_break();d.add_paragraph(title,'Heading 1')
def table(d,heads,rows,widths=None):
 t=d.add_table(rows=1,cols=len(heads));t.autofit=False
 for i,x in enumerate(heads):t.rows[0].cells[i].text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for ri,r in enumerate(t.rows):
  for i,c in enumerate(r.cells):
   if widths:c.width=Cm(widths[i])
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
   pr.append(b)
   if ri==0:
    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7EEF4');pr.append(fill)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(4);pp.paragraph_format.line_spacing=1.1
    for run in pp.runs:run.font.size=Pt(9);run.bold=ri==0
  trpr=r._tr.get_or_add_trPr()
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
 d.add_paragraph('')
 return t

# Narrow edit of a retained application: only selected table cells change.
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};W='{'+ns['w']+'}'
texts={
5:'面向可靠图像编辑的诊断式评价方法研究',
10:'修订定位：本副本将原“评价与统计”职责拓展为可独立开展的诊断式评价研究，保留团队评分细则、校准、盲评组织与统计职责。属导师讨论稿，不代表队友已同意新分工；本表仍为原 ESRT 模板，不冒充学校大研正式表格。\n研究问题：面对编辑中的可疑变化，评价器能否学习诊断成立、撤销与保持的条件，并自主选择下一步检查的区域、属性和工具，区分合理关联变化与非预期损伤？\n研究内容：（1）构建原图—指令—结果的区域及属性诊断表示，分别记录完成度、内容保持、局部质量与不确定性。（2）用对应修正、无关改善及新增损伤对照，训练轻量诊断模块，要求修好对应错误后撤销诊断，无关美化不改变判断。（3）训练检查价值模块，在有限工具集合内学习下一步检查哪里、检查什么及何时停止。（4）以真实结果、未见编辑器与工具故障测试检验泛化，采用独立人工评价与等成本对照。\n独立实施：先使用许可允许的公开编辑结果或固定工具少量生成样例，不以团队反事实动作表或决策器完成为前提。初期冻结基础模型，训练小头，必要时参数高效适配；不训练编辑器。先做20组试标，再根据学习曲线扩充数据，规模与算力待导师确认。\n团队衔接：向队友提供初始结果的结构化诊断、证据和检查成本，由队友预测并选择编辑动作。可辅助训练标签质量审核，但不把本模型同时作为最终成绩的唯一裁判。诊断检查策略与实际改图策略分别归属，原六动作保持由团队实现。',
11:'独立成果：①诊断与对照标注协议及可发布的数据清单；②可检验诊断模型与下一步检查选择器；③真实性、干扰稳定性、跨编辑器和成本实验；④可复现代码及个人研究报告。\n团队成果：评分与校准手册、诊断接口、盲评组织和统计附录，以及保持同一决策算法的输入消融实验。论文投稿是争取目标，不作为必然产出。\n候选贡献：诊断响应监督与面向疑点的证据价值学习；定位、工具调用和反馈修复本身已有先例，须与近邻工作充分比较。导师指导和组外复核均如实署名。',
12:'建议按30个有效工作周推进，每周平均12小时，其中个人研究约10小时、团队职责约2小时；另留约4个日历周应对考试与排队。\nW1–4：文献、资源核查、强基线与20组试标；W5–9：诊断协议、数据划分与工具缓存；W10–15：对照诊断训练；W16–21：区域与检查选择策略；W22–26：独立泛化、故障与等成本验证；W27–30：封存测试、报告、团队接口及复现。\n个人独立工作预算260–340小时，团队额外职责40–60小时，合计300–400小时；基准计划为300+60=360小时，不含组外评分者工时。若长期每周仅8小时，总历时约38–50个有效周。具体日期按导师启动与学校通知调整。',
13:'预算改为待资源核实的规划区间，不沿用原1800元作为充分经费承诺：独立课题算力/API暂预留800–1800元，独立数据复核400–800元，存储与重跑200–400元，合计1400–3000元。均为非报价的预算占位，按20组试跑实测修订；实验室提供资源时记录实物支持。团队原大规模盲评与动作生成另列团队预算，不重复报销，不把900元视为必然足够。',
14:'已有基础：团队研究文档、评价协议草案与文献调研；可考虑公开编辑结果和开源评分、检测、分割及OCR组件。实际GPU、API额度和导师均待确认，不声明已获实验室支持。每周可投入8–15小时，目前约一位同学可协助标注；团队最终双人组外盲评还需协调第二位组外人员。本人有算法实践经历，但模型适配、统计与可靠性实验需在导师指导下补齐。'
}
with ZipFile(sources[0]) as z:
 xml=E.fromstring(z.read('word/document.xml'));rows=xml.findall('.//w:tbl',ns)[0].findall('w:tr',ns)
 for idx,txt in texts.items():
  tc=rows[idx].findall('w:tc',ns)[1];old=tc.find('w:p',ns);pp=deepcopy(old.find('w:pPr',ns)) if old is not None and old.find('w:pPr',ns) is not None else None
  rp=old.find('.//w:rPr',ns) if old is not None else None
  for child in list(tc):
   if child.tag!=W+'tcPr':tc.remove(child)
  for line in txt.split('\n'):
   el=E.SubElement(tc,W+'p')
   if pp is not None:el.append(deepcopy(pp))
   r=E.SubElement(el,W+'r')
   if rp is not None:r.append(deepcopy(rp))
   E.SubElement(r,W+'t').text=line
 app=OUT/'ESRT_立项申请表_huangtianzhi_诊断研究修订副本_20261028.docx'
 with ZipFile(app,'w') as dest:
  for entry in z.infolist():dest.writestr(entry,E.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True) if entry.filename=='word/document.xml' else z.read(entry.filename))

d=setup('黄天志个人独立研究方案')
p(d,'面向可靠图像编辑的诊断式评价方法研究\n导师讨论稿 · 个人工作与团队职责分别核算')
h(d,'定位与独立性')
p(d,'研究一个会主动查证的图像编辑质检员：它理解指令、提出疑点，学习下一步检查哪里与什么属性，再根据证据保留或推翻诊断。独立课题可使用公开结果或固定工具生成数据，不等待团队动作表和决策器。独立是研究问题和实验链条可单独成立，不排斥导师指导、组外复核或共同署名。')
h(d,'拟解决的两个核心问题')
p(d,'一是可检验诊断学习。利用“对应错误被修正”和“无关内容被改善”等对照，学习诊断应何时撤销、保持或更新。二是主动核验学习。根据当前疑点与证据，选择区域、属性及检查方式，并决定何时停止。对合理关联变化的辨别是关键挑战，修复收益是应用扩展。')
h(d,'起步任务与边界')
p(d,'先覆盖改色与局部属性编辑中可核实的文字、对象和边界保持；再加入少量反光、材质等合理关联变化。起步场景简单是为验证机制，不把课题缩成固定分类器。第一版不开发通用修复器、不研究全部物理规律、不承担队友的编辑动作选择。')
table(d,['条件','对方案的影响'],[
['每周8–15小时','按平均12小时做基准，考试期降低强度并顺延'],['算力与API未知','先冻结骨干、缓存证据、小头训练，再决定适配'],['约1位标注助手','个人承担主标；独立测试双评，团队组外盲评另协调'],['原团队职责仍存在','单独预留40–60小时，不能再全部算作独立研究时间']],[4,12])
p(d,'阶段成果可以是可信的负结果与边界分析；不保证首次、投稿等级或预定性能提升。正式学校申报需使用当年模板，ESRT副本仅为原申请的讨论修订。')

page(d,'技术路线与个人交付物')
table(d,['模块','个人工作','复用基础'],[
['要求与诊断','定义对象 属性 要求 证据 不确定性字段','冻结多模态骨干提出候选'],['响应学习','真实性 修正撤销 无关稳定 新损伤监督','缓存局部图像与表示'],['主动核验','区域×属性×工具候选价值学习及停止策略','裁剪 OCR 检测 分割'],['验证','独立真值 成本 泛化 故障测试','强评价模型与固定检查基线'],['输出','纠错及保护条件与版本化接口','固定工具仅用于应用验证']],[3,8,5])
h(d,'学习实施')
p(d,'诊断头学习当前错误是否真实，并约束对照间的响应；真实性监督防止“永远没有错”的退化解。检查头在训练时利用候选检查前后相对独立标签的诊断损失下降与实测成本学习排序，推理时不能访问标签。先监督学习，必要时再研究短序列策略，不默认采用强化学习。')
p(d,'区域候选同时包括对象框、差异区域和覆盖性网格，避免永远只检查已怀疑位置。OCR不可读、对齐失败和工具缺失均有独立状态，不可编码成“无损伤”。小头不等于低总成本，骨干推理与工具调用均记录。')
h(d,'独立完成的验收产物')
p(d,'数据说明与拆分清单；诊断标注协议；对照诊断模型；下一步检查选择器；至少一个真实数据和一个留出编辑器实验；固定清单、直接工具路由与完整方法消融；个人报告和复现脚本。团队联合性能提升不作为个人课题完成的唯一条件。')
h(d,'扩展储备')
p(d,'所有28个idea保留在研究全景文档。复杂变化依赖图、在线图像干预、多轮历史追踪、反馈收益与保护条件学习为扩展；优先保留对照诊断、主动检查、不确定性与真实泛化。模块更多不自动带来鲁棒性。')

page(d,'数据设计与实验可信度')
p(d,'先用20组样例试标和计时，检查人类是否能判断合理变化与真实损伤。首轮可规划200组训练开发（例如160训练、40开发）与100组独立测试，每组约3–4种有效变化；组内变体并非独立样本，规模由试跑和学习曲线调整，不宣称已足够。')
p(d,'训练可用自动候选与伪标签，但人工核验修改是否真的只改变目标因素。按原图先划分再构造变体；公开基准作为测试时不得同时用于训练。个人数据与团队封存测试做ID或图像哈希排重；团队开发样例若训练评价器，就不能继续充当该评价器的独立校准集。')
table(d,['比较','回答的问题'],[
['直接评价 vs 固定清单','流程与额外观察本身有多少收益'],['普通监督 vs 对照监督','可检验诊断学习是否有效'],['固定/随机/基础模型路由 vs 学习选择','是否学会了有价值的下一步检查'],['全量检查 vs 主动检查','成本优势及性能上限'],['真实与合成分报 留出编辑器','是否仅识别合成痕迹或编辑器风格']],[7,9])
p(d,'报告误报漏报、修正撤销率、无关变化错误翻转、新损伤检出、合理变化误罚、风险覆盖曲线及准确性成本曲线。所有策略使用相同工具与信息权限。按原图分组估计置信区间，不把4个变体当4张独立原图。最终真值不能只来自自己的评价器。')
h(d,'标注预算')
p(d,'本人标注、规则修订与仲裁预计40–60小时，包含在个人数据阶段工时内，不另加一次。助手预计10–20小时另列。粗算：200组主标每组4–8分钟约13–27小时；25%复核约3–7小时；100组双评约13–27合计人时，再加规范与仲裁。两者是不同口径，应据试标统一更新。复杂核验轨迹可能提高时间。')
p(d,'一个人工作一小时为一人时。组外双评是团队最终结论要求；个人方法开发者可以参与个人测试标注，但需如实披露并由独立复核者核对，不能称双人组外盲评。若要求该标准，必须再落实第二位组外人员。')

page(d,'工作量与推进时间线')
table(d,['阶段','周次','个人小时','团队小时','交付与判据'],[
['定题与试跑','1–4','40','8','20组试标 资源清单 强基线失败案例'],['数据与基础实现','5–9','50','10','拆分冻结 工具缓存 标注协议'],['诊断学习','10–15','60','12','训练模型 真实开发集响应实验'],['主动核验','16–21','60','12','区域与检查选择 等预算消融'],['独立验证','22–26','50','10','泛化 故障 成本及封存准备'],['总结与对接','27–30','40','8','个人报告 接口复现 联合实验或待接入包'],['合计','30周','300','60','360小时 每周平均12小时']],[3,1.8,2,2,7.2])
p(d,'总量估算：个人260–340小时、团队额外40–60小时，合计300–400小时。前表为360小时基准，不是对每个阶段的精确预测。标注包含于个人工作，不重复累计；助手和其他组外评分者工时不算你的个人小时。GPU等待不按持续人工工作计时。')
table(d,['每周总投入','300–400小时所需有效周数'],[['8小时','约38–50周'],['12小时','约25–34周'],['15小时','约20–27周']],[6,10])
p(d,'若2026年11月2日启动，30个连续有效周到2027年5月30日；另留约4个日历周处理考试、节假日或GPU排队，目标可放至6月下旬。仅为排期示例，不代表学校开题或结题期限。若只能每周8小时，应延长周期或取消扩展，不能按12小时计划硬赶。')
p(d,'每周基准：5小时实现与训练、3小时数据或实验、2小时阅读与记录、2小时团队任务。阶段不同可调整。每两周提交一个可复现结果给导师；每四周重新估算剩余工作。团队动作表延迟时，继续个人数据和诊断实验，不空等。')

page(d,'资源 里程碑与导师讨论')
h(d,'需要尽早确认的资源')
p(d,'确认可用GPU显存、预约时段与能否长期缓存；API是否有经费；数据及模型许可；标注助手可投入时数；导师沟通频率。先测20组的耗时、失败率和显存，再选骨干与训练方式。没有GPU时可先做公开结果和离线特征实验，但不能隐瞒特征获取成本。')
p(d,'经费暂按1400–3000元为个人预算占位：算力/API 800–1800、复核400–800、存储重跑200–400。不是报价，也不表示学校会批准。按试跑换算机时与调用费用；团队4.3万次级生成与最终盲评另列，不能塞进这份个人预算。')
h(d,'继续或调整的判据')
p(d,'W4：若人工共识不足，先修定义；若强基线已足够好，重新定位真难例。W15：若合成收益不迁移，调整数据和表示，不盲目补标。W21：若主动策略不胜固定清单，如实保留诊断学习贡献并分析成本边界。封存前冻结配置；若工具质量制约方法，人工可靠证据仅作上限诊断，不能作为部署成绩。')
h(d,'需要导师判断的三件事')
p(d,'一是可检验诊断与证据选择的具体监督是否具有实质差异；二是两类最值得研究且可获得人工共识的失败场景；三是资源是否支持一学期以上的个人训练与验证。独立课题不承诺全部想法同时落地，核心成果应有代码、数据依据和可信实验。')
h(d,'主要近邻与差异假设')
p(d,'EditProbe已有中间层评分适配，SpatialReward已有区域核验，RewardHarness已有工具技能选择，ReasonEdit已有解释质量学习，R³-Refiner已有反馈修复。拟新增的是诊断在受控变化下的响应监督，以及由这些条件驱动的证据价值学习。差异需同工具、同预算实验确立，不能以改名或接入开源替代。')
for txt in ['EditProbe https://arxiv.org/abs/2603.19775','SpatialReward https://arxiv.org/abs/2602.07458','RewardHarness https://arxiv.org/abs/2605.08703','ReasonEdit https://arxiv.org/abs/2605.07477','R³-Refiner https://arxiv.org/abs/2605.19639']:
 pp=d.add_paragraph(txt);pp.paragraph_format.space_after=Pt(2)
 for r in pp.runs:r.font.size=Pt(8)
personal=OUT/'htz_个人独立研究方案与时间线_20261028.docx';d.save(personal)

d=setup('黄天志研究与团队接入方案')
p(d,'基于指代式图像编辑空间控制决策的分工修订建议\n讨论副本 · 原团队文档保持不变 · 待三人及导师确认')
h(d,'原分工与本次调整')
p(d,'原文件表8将孙彧负责的预测决策、何劲负责的控制机制与黄天志负责的评价统计分开。本次保留前两位的动作实现、收益预测和预算策略边界；黄天志在原有协议校准与统计之上，增加独立的可检验诊断和主动检查学习。下列是建议，不代表队友已经同意。')
table(d,['责任','旧安排','建议安排'],[['黄天志','评分细则 校准 盲评 统计','上述职责保留；新增诊断模型与检查选择；不独占全部人工评分'],['孙彧','动作标签特征 预测器 策略','接收版本化诊断特征，仍负责编辑动作收益与选择'],['何劲','编辑器 六动作 扰动 动作表','仍负责编辑控制和生成；按接口提供原图初始结果与动作记录']],[3,5,8])
h(d,'并行工作流')
p(d,'第一条线：何劲核验编辑器并建立动作表。第二条线：黄天志用公开结果或少量固定生成独立开发诊断器。第三条线：孙彧先用基础状态与占位接口开发预测器。三条线共享样本ID、数据划分及成本定义，在各自就绪后接入，不把五步文字理解为必须串行等待。')
p(d,'W1–4先交付诊断字段和最简固定检查实现；W5–9交付工具输出；W10–15交付训练诊断；W16–21交付主动核验。团队可随时使用最近稳定版本；联合测试前固定版本，不能边看测试结果边更新评价器。')
h(d,'原文件对应位置')
p(d,'第3.1节流程第③步扩展为诊断与可信评价并行线；第3.3节允许加入初始结果诊断特征；第3.5及4.5节保留独立校准和双人组外盲评；第6.2节更新黄天志交付物。原六动作、训练开发测试划分与其他成员研究问题不在此副本中直接改写。')

page(d,'接口与公平实验')
table(d,['接口项','要求'],[['输入','sample_id source_image instruction initial_output 可选原始掩码'],['诊断输出','object region attribute requirement status severity uncertainty evidence'],['执行记录','inspection_cost 工具失败日志 evaluator_version schema_version'],['未知处理','明确不可判定 不把空值当正常'],['禁止输入','待执行动作的真实结果 测试标签 参照答案等未来信息']],[4,12])
p(d,'主要接入点是初始结果y0产生后、实际编辑动作执行前。评价器只选择检查动作。队友接收诊断，预测收缩、扩张、精修、重定位、重采样或保持的价值。原动作从原图执行，不把它们暗改成对y0的连续修复；如需要反馈驱动编辑另设扩展协议。')
h(d,'成本与标签隔离')
p(d,'总预算必须包含诊断、分割、OCR、编辑、重试等成本。同预算比较中，增加检查可能减少可用编辑次数，应报告实际结果。用于训练动作预测器的自动标签可经诊断器审核，但只允许训练数据；低一致性标签应限制使用，不能认为加一个置信度特征就消除了偏差。')
p(d,'评价器若用于个人训练的数据与团队训练数据重合，需要记录来源并防止同源变体跨集合。团队开发集用于训练后不能继续作为评价器独立校准集；最终测试仍保持封存。必要时采用训练内交叉拟合，禁止从测试反向生成训练标签。')
h(d,'建议联合消融')
p(d,'保持队友的算法、训练预算和输入维度适配规则，比较基础状态、标量评价、固定检查诊断、主动检查诊断。若重训预测器，说明各组同配置；若只换输入而不重训，需要证明接口与分布适配。最终结果由独立人工或外部评价核实，不能让新评价器给自己参与优化的系统独自打分。')
p(d,'联合实验依赖团队进度，但个人课题的诊断准确性、响应一致性和检查成本实验可以独立完成。团队暂未就绪时交付模拟输入测试、接口包与个人结果，不冒充已获得联合性能提升。')

page(d,'人力预算与协作约定')
p(d,'原方案开发校准约300例双评；最终测试720任务组，每组1–3个互异结果，双人组外评价。合计2040–4920次评分，按1–2分钟约34–164合计人时，另加仲裁和培训。该数字来自原文，不是黄天志个人标注要求；100图掩码精修另属团队数据工作。')
table(d,['工时归属','建议口径'],[['黄天志独立研究','260–340小时；基准300；含个人标注与实验'],['黄天志团队额外职责','40–60小时；基准60；协议协商 校准配额 接口 盲评组织 统计'],['其他标注者','独立复核10–20小时为初估；团队双人组外盲评另计'],['共享工作','协议 工具或数据清洗能复用的部分只记一次；工时台账标记个人/团队/共享']],[5,11])
p(d,'目前仅约一位同学协助标注，不满足原方案两位组外评分者的资源要求。开发双标三人轮换；组外招募、评分费用与争议仲裁由团队共同协调。黄天志负责组织不等于代替所有评分者。若最终评分量超资源，三人需事先讨论缩减测试覆盖和相应结论，不能解封后选择性删样本。')
h(d,'建议预先确认')
p(d,'确认诊断与实际编辑动作的职责边界；确认训练开发测试的共享规则；确认组外人员与费用；确认每两周接口联调；确认联合实验以稳定版本为准；确认个人研究报告、共享代码、数据贡献与共同论文署名如实记录。')
h(d,'实施范围说明')
p(d,'本文件是原团队方案的接入修订说明，未覆盖或替代原团队文档，也未未经讨论改变其他成员的任务、主检验或预算。与个人方案、ESRT修订副本共同构成导师讨论材料；正式采用前由团队确认。')
team=OUT/'htz_团队接入与分工修订说明_20261028.docx';d.save(team)
assert all(hashlib.sha256(p.read_bytes()).hexdigest()==hashes[p.name] for p in sources)
with ZipFile(sources[0]) as a,ZipFile(app) as b:
 changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
 assert changed==['word/document.xml'],changed
(QA/'verification.json').write_text(json.dumps({'originals_unchanged':True,'application_changed_parts':changed,'outputs':[str(x) for x in [app,personal,team]]},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'outputs':[str(x) for x in [app,personal,team]],'originals_unchanged':True},ensure_ascii=False))
