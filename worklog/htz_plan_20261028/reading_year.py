import ast,json
from pathlib import Path
BASE=Path(__file__).resolve().parent
tree=ast.parse((BASE/'build.py').read_text(encoding='utf-8'))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n,(ast.Import,ast.ImportFrom,ast.FunctionDef))],type_ignores=[]),'helpers','exec'))
OUT=BASE.parents[1]/'docs'
from docx.opc.constants import RELATIONSHIP_TYPE as RT
def link(d,label,url):
 pp=d.add_paragraph();pp.paragraph_format.space_after=Pt(5)
 a=OxmlElement('w:hyperlink');a.set(qn('r:id'),d.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'245980');pr.append(c);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);a.append(r);pp._p.append(a)

# title, author, date/status, URL, contribution, reading focus, research question
A=[
('EditProbe','Evaluating Image Editing with LLMs: A Comprehensive Benchmark and Intermediate-Layer Probing Approach','Shiqi Gao 等','2026-03；arXiv v2 03-25；本轮出版社页面受限，期刊状态不作确认','https://arxiv.org/abs/2603.19775v2','以TIEdit人工评分与中间层表示学习编辑质量，采用参数高效适配；不是单纯冻结模型接一个评分头。','方法IV的层选择、AdaLoRA、实验切分及消融；先弄清输入、标签和预测目标。','同一骨干上，对照诊断监督能否比普通分数监督更可靠？阅读产物：画出最小可复现基线。'),
('ImageDoctor','ImageDoctor: Diagnosing Text-to-Image Generation via Grounded Image Reasoning','Yuxiang Guo 等','ICLR 2026；正式论文集已核实','https://proceedings.iclr.cc/paper_files/paper/2026/hash/deb0e85779f2f003b10528de72b1ebf1-Abstract-Conference.html','以定位、推理和预测输出多维分数与缺陷热图，并研究稠密奖励。已说明“诊断＋定位”有直接先例。','模型结构、热图监督、定位与推理消融；区分文生图评价与原图条件编辑评价。','热图正确是否意味着诊断能对修正与无关变化作出正确响应？产物：列出新增响应测试。'),
('SpatialReward','SpatialReward: Bridging the Perception Gap in Online RL for Image Editing via Explicit Spatial Reasoning','Yancheng Long 等','arXiv v4 2026-05-13；作者注释称ICML 2026录用，会议页未独立复核','https://arxiv.org/abs/2602.07458v4','预测编辑区域以锚定空间推理与跨图核验，训练专用评价模型。不能把“学会看哪里”泛称首创。','Think-with-Boxes、数据构造、空间定位消融及成本；以v4正文核对旧版差异。','一次定位与逐步选择新增证据的差别能否实测？产物：设计同区域候选下的固定与主动对照。'),
('ReasonEdit','ReasonEdit: Towards Interpretable Image Editing Evaluation via Reinforcement Learning','Honghua Chen 等','2026-05-08；arXiv预印本，正式状态未核实','https://arxiv.org/abs/2605.07477','人工评价解释文本的逻辑性、准确性与有用性，训练RE-Reward并优化解释式评价器。','解释标注协议、RE-Reward、训练目标和消融；注意同名CVPR编辑生成论文不是本篇。','解释看似有用与实际核验有效有何区别？产物：将文本质量监督与对照响应监督并列。'),
('RewardHarness','RewardHarness: Learning Human Preferences for Image Editing with Only 100 Demonstrations','Yuxuan Zhang 等','作者页面称COLM 2026；OpenReview验证页受限；arXiv v2更新于2026-10-12','https://arxiv.org/abs/2605.08703v2','通过偏好示例演化工具技能库，选择上下文指导冻结模型评价。预印本题名为Self-Evolving Agentic Post-Training。','工具是实际执行器还是过程描述、路由时机、演化验证与成本；会议版和v2不能混算指标。','我们的逐步证据价值学习比库路由多了什么？产物：画状态、动作、监督的逐项差异表。'),
('R³-Refiner','Benchmarking and Evolving Reason-Reflect-Rectify for Reflective Visual Generation','Junjie Wang 等','ICML 2026；正式论文集核实','https://proceedings.mlr.press/v306/wang26fe.html','学习反思和纠错指令，交由外部编辑器执行；评价诊断与纠错两个阶段，并研究训练编辑器迁移。','两阶段奖励、冻结编辑器、反思与纠错指标、迁移附录；区分训练自评信号和最终外部验证。','如何区分诊断真伪与修复器执行能力？产物：固定编辑器反馈实验的混杂清单。'),
('CheckDistill','CheckDistill: Adaptive Checklist Distillation for Efficient Distributional Reward Modeling in Image Editing','Yaqi Cai 等','2026-10-19；arXiv v1预印本','https://arxiv.org/abs/2610.26740v1','将指令分成可验证检查项，通过特权教师蒸馏给直接输出分数分布的学生。','检查项构造、可靠性加权蒸馏、推理时输入与成本；不要误以为部署时仍生成长清单。','主动执行检查与把检查知识蒸馏进评分器谁更划算？产物：列出需要外部新证据的样例。'),
('EditJudgeBias','Do MLLM Judges Judge the Edit? Auditing Bias in Image Editing Evaluation with Verified Quality Preservation','Yuan Huang 等','2026-10-01；arXiv v1预印本','https://arxiv.org/abs/2610.01670v1','在核实编辑质量保持的条件下审计无关线索引起的判分偏移，考虑重查询噪声和偏好稳定性。','干预有效性审核、零剂量对照、重查询噪声、人工检查；不能把所有分数变化都归为偏差。','我们的无关变化是否真的不影响质量？产物：对照有效性审核表与噪声基线。'),
('EVR','Evaluation-Verification Reward for Consistent Multi-Reference Image Editing','Yingmao Miao 等','2026-07首发；v2 2026-08-05；正式状态未核实','https://arxiv.org/abs/2607.29025v2','多参考编辑中由评价器提出候选假设，验证器基于视觉证据接受或拒绝，聚合成奖励。','假设产生、证据验证与奖励聚合；重点查是否主动选择新观察，还是验证已有输入。','“提出假设再验证”已被覆盖，我们新增的学习状态与收益在哪里？产物：最接近方法对照。'),
('EditHF','EditHF-1M: A Million-Scale Rich Human Preference Feedback for Image Editing','Zitong Xu 等','2026-03-16；arXiv预印本，正式状态未核实','https://arxiv.org/abs/2603.14916','大规模多维人类反馈训练评价模型并用于编辑优化。偏好对规模不等于逐对独立人工劳动量。','标注流程、偏好构造、训练切分、跨数据集测试和评价指标。','少量诊断对照是否在等人工成本下更有效？产物：标注成本与学习曲线方案。')
]
B=[
('EditScore','EditScore: Unlocking Online RL for Image Editing via High-Fidelity Reward Modeling','Xin Luo 等','v3 2026-02；作者注释ICLR 2026录用','https://arxiv.org/abs/2509.23909v3','专用编辑评分、基准与在线RL应用。','读数据过滤、推理和自集成；选一个可运行版本作强基线，不复现全部RL。'),
('EditReward','EditReward: A Human-Aligned Reward Model for Instruction-Guided Image Editing','Keming Wu 等','v2 2026-02；作者注释ICLR 2026录用','https://arxiv.org/abs/2509.26346v2','人类偏好训练奖励模型，并以筛选编辑训练数据验证用途。','读偏好协议及筛选实验；与EditScore同名基准须核对作者、数据版本与来源。'),
('PVC-Judge','GEditBench v2: A Human-Aligned Benchmark for General Image Editing','Zhangqi Jiang 等','2026-03-30；arXiv预印本','https://arxiv.org/abs/2603.28547','提出视觉一致性成对评价器PVC-Judge与区域解耦偏好合成。','读一致性定义与合成流程；“区域保持评价”有直接近邻，核查错位与伪标签偏差。'),
('Edit-Compass','Edit-Compass & EditReward-Compass: A Unified Benchmark for Image Editing and Reward Modeling','Xuehai Bai 等','v2 2026-10-05；arXiv预印本','https://arxiv.org/abs/2605.13062v2','统一编辑能力与奖励模型评价，强调细粒度规则和实际偏好场景。','读评分规则及候选对构造；据此设计难例，不照搬整套基准。'),
('SQUARE-Bench','Evaluating the Evaluators: Diagnosing Large Multimodal Models for AI-Generated Image Assessment','Yu Zhao 等','2026-09-29；arXiv预印本','https://arxiv.org/abs/2609.37576','评价多个维度，并用维度诊断指导固定编辑器；部分提升伴随视觉质量权衡。','读分类体系及闭环实验；改好某维度可能引入新损伤，适合保护条件研究。'),
('UMM-Reflection','Learning Native Reflection in Unified Models with Interleaved Reinforcement Learning','Yijia Fan 等','v2 2026-10-20；arXiv预印本','https://arxiv.org/abs/2609.35767v2','以完整反思修正轨迹联合训练统一模型的文本与图像生成。','读轨迹奖励和信用分配；你的外置评价器不必复现昂贵联合训练。'),
('Inline Critic','Inline Critic Steers Image Editing','Weitai Kang 等','2026-05-12；本清单按arXiv版本，正式状态未核实','https://arxiv.org/abs/2605.12724','在编辑模型中间层用可学习critic token引导后续生成。','读探针与分阶段训练；区别内部生成控制和外部诊断，避免无意接管编辑器。'),
('PreActIR','Look Before You Restore: Learning to Predict Before Acting for Agentic Image Restoration','Xu Zhang 等','2026-10-14；arXiv注明Under review','https://arxiv.org/abs/2610.20384','执行前预测工具收益、风险与成本，验证后提交或回滚。','读候选空间审计与等成本实验；对应队友的编辑动作选择，你选择的是检查动作。'),
('ImagenWorld','ImagenWorld: Stress-Testing Image Generation Models with Explainable Human Evaluation on Open-ended Real-World Tasks','Samin Mahdizadeh Sani 等','ICLR 2026；正式论文集','https://proceedings.iclr.cc/paper_files/paper/2026/hash/458fa8ee331566383d8e74bdb647f829-Abstract-Conference.html','真实开放任务与局部错误的人类解释，检验生成及编辑模型。','读人工错误标签及域划分；整体排序好不代表细粒度归因正确。'),
('DLEBench','DLEBench: Evaluating Small-scale Object Editing Ability for Instruction-based Image Editing Model','Shibo Hong 等','ICML 2026；正式论文集','https://proceedings.mlr.press/v306/hong26a.html','小目标编辑基准，工具驱动与oracle引导评价。','读两种评价模式；人工可靠区域只能估计上限，不能当部署效果。'),
('EditInspector','EditInspector: A Benchmark for Evaluation of Text-Guided Image Edits','Ron Yosef 等','ACL 2025；正式论文集；重要前作','https://aclanthology.org/2025.acl-long.1428/','人工模板覆盖编辑验证、伪影和差异描述，并提出相关方法。','读标注模板及幻觉分析；复用标注结构，不能称首次细粒度诊断。'),
('BPM','Balancing Preservation and Modification: A Region and Semantic Aware Metric for Instruction-Based Image Editing','Zhuoying Li 等','ICML 2025；正式论文集；重要前作','https://proceedings.mlr.press/v267/li25dh.html','区域与语义两级评判完成与保持，定位还能辅助编辑。','读区域划分与检查流程；分离完成度和保持性本身已被覆盖。'),
('AugCLIP','Preserve or Modify? Context-Aware Evaluation for Balancing Preservation and Modification in Text-Guided Image Editing','Yoonjeon Kim 等','CVPR 2025；正式论文；早期预印本为2024','https://arxiv.org/abs/2410.11374','根据上下文构造保持与修改的CLIP空间目标。','读属性提取与理想编辑表示；理解“允许变化”不是只看掩码外差异。'),
('EdiVal-Agent','EdiVal-Agent: An Object-Centric Framework for Automated, Fine-Grained Evaluation of Multi-Turn Editing','Tianyu Chen 等','v3 2026-03-17；正式状态未核实','https://arxiv.org/abs/2509.13399v3','对象池随编辑更新，结合检测、局部语义核验与保持评价。','读对象更新和检测错误；多轮追踪可后置，工具失败不能视作真实图像错误。'),
('PELC','Your Latent Mask is Wrong: Pixel-Equivalent Latent Compositing for Diffusion Models','Rowan Bradbury 与 Dazhi Zhong','2025-12-04预印本；本轮未核实正式会议状态','https://arxiv.org/abs/2512.05198','用像素等价原则和DecFormer处理潜空间混合造成的边界及全局伪影。','读失效机制；帮助构造真实边界损伤，主要属控制队友参考，不要求个人复现。'),
('MDC','Model Diagnosis and Correction via Linguistic and Implicit Attribute Editing','Xuanbai Chen 等','CVPR 2025；正式论文；本轮新补近邻','https://openaccess.thecvf.com/content/CVPR2025/papers/Chen_Model_Diagnosis_and_Correction_via_Linguistic_and_Implicit_Attribute_Editing_CVPR_2025_paper.pdf','从视觉模型错误发现属性，通过编辑做反事实假设检验，并用样本改善识别模型。','重点读假设验证；“编辑属性检验模型判断”已有先例，差异需落到编辑诊断响应与主动观察。'),
('Meta-CoT','Meta-CoT: Enhancing Granularity and Generalization in Image Editing','Shiyi Zhang 等','CVPR 2026；CVF页面及论文核实','https://arxiv.org/abs/2604.24625','分解任务、目标与理解能力，利用元任务和推理编辑一致性奖励训练编辑器。','读分解方式及一致性定义；你的诊断表示可借鉴，但不能仅凭拆任务宣称创新。'),
('Rubric-CEPR','Rubric-CEPR: Self-Evolving Image Editing via Reward-Verified Self-Distillation','Ritesh Thawkar 等','v2 2026-10-09；arXiv预印本；本轮新补','https://arxiv.org/abs/2610.12469v2','从编辑器自身表示构造完成、旧状态移除与保持检查，筛选候选用于适配。','读不可补偿门控和自评风险；综合评分不能让美观抵消严重失败。'),
('HYPE-EDIT-1','HYPE-EDIT-1: Benchmark for Measuring Reliability in Frontier Image Editing Models','Wing Chan 与 Richard Allen','2026-01-25；arXiv预印本；本轮新补','https://arxiv.org/abs/2602.00105','以重复尝试通过率、重试和人工复核代价衡量实际可靠性。','读成本口径；你的主动核验成本应包含人工与失败重试，但不要照搬其具体价格。'),
('DACR与ARE','Adaptive Robust Estimator for Multi-Agent Reinforcement Learning','Zhongyi Li 等','2026-03-23；本清单按arXiv，正式状态未核实','https://arxiv.org/abs/2603.21574','分解回答、批评与改写并研究边际贡献及噪声奖励下的稳健优化。','只读交互协议与奖励归因；跨领域提醒，边际反馈收益不是全新一般概念。')
]
assert len(A)==10 and len(B)==20

d=setup('诊断式图像编辑评价文献阅读指南')
p(d,'必读10篇与需要了解20篇\n核查截止 2026年10月28日 · 附初步综述与12周阅读路径')
h(d,'选择原则与证据边界')
p(d,'优先纳入2026年新成果与截至检索日可确认的更新版本。必读10篇以当前课题的直接近邻为主；另20篇覆盖基准、保持评价、反馈修复、工具与反事实方法，其中保留少量2025年重要前作。最新不等于最好，先读与你的核心假设最接近的工作。')
p(d,'正式论文集、作者录用声明和预印本分别标注。未核实正式发表不等于没有发表；预印本也影响创新判断。本轮核查题名、版本、摘要与关键方法线索，结合前轮正文阅读，不宣称已经逐篇精读30篇全文或复现结果。出版社与OpenReview部分页面受限时明确注明，不据搜索摘要推断录用。')
h(d,'研究主线')
p(d,'学习诊断成立、撤销和保持的条件，再学习当前疑点下一步最值得检查的区域、属性和工具。阅读时特别检查四个问题：模型是否获得新证据；是否真正训练检查策略；监督来自人工、伪标签还是自评；改进是否在等成本与独立测试下成立。')
h(d,'版本与同名陷阱')
p(d,'RewardHarness会议版与预印本题名不同，勿混用表格数值；ReasonEdit有评价和编辑生成的同名工作，本清单A04是评价论文；SpatialReward也有同名文生图空间评分工作，本清单A03是图像编辑论文；PVC-Judge是GEditBench v2中的方法名，不是该论文完整题名。')
h(d,'工时预算')
p(d,'必读每篇约2.5–3.5小时，了解每篇约0.5–0.75小时，另用约8–10小时整理矩阵和综述、约5–8小时追踪版本与补读，约48–68小时。年度基准计60小时，已包括在总工作量中。基础薄弱时需增加补课时间；不要求每篇复现。')
p(d,'建议首先阅读A01、A02建立概念，再读A03–A09确认近邻边界。A10帮助判断数据规模和标注成本。B16虽为2025年前作，涉及反事实假设检验，应在提出“新范式”前提前阅读。')

md=['# 诊断式图像编辑评价文献阅读指南','核查截止2026-10-28。10篇必读＋20篇了解。状态和版本见各条；非穷尽查新。']
for start in range(0,10,2):
 page(d,f'必读文献 A{start+1:02d}至A{start+2:02d}')
 for i in range(start,start+2):
  short,title,author,status,url,claim,focus,q=A[i]
  h(d,f'A{i+1:02d} {short}');p(d,title);p(d,author+'｜'+status);link(d,'论文及版本入口',url)
  p(d,'核心内容：'+claim);p(d,'重点阅读：'+focus);p(d,'带着问题读：'+q)
  md.extend([f'## A{i+1:02d} {title}',author+'｜'+status,f'[论文入口]({url})','核心内容：'+claim,'重点阅读：'+focus,'研究问题：'+q])
for start in range(0,20,4):
 page(d,f'需要了解 B{start+1:02d}至B{start+4:02d}')
 for i in range(start,start+4):
  short,title,author,status,url,claim,focus=B[i]
  h(d,f'B{i+1:02d} {short}');pp=d.add_paragraph(title);pp.paragraph_format.space_after=Pt(3)
  pp=d.add_paragraph(author+'｜'+status)
  for r in pp.runs:r.font.size=Pt(9)
  link(d,'论文入口',url);p(d,claim+' '+focus)
  md.extend([f'## B{i+1:02d} {title}',author+'｜'+status,f'[论文入口]({url})',claim,focus])

page(d,'初步综述与研究缺口')
review=[
('从整体分数走向专用评价','EditScore、EditReward、EditHF与EditProbe分别从奖励模型、偏好数据和表示学习推进编辑评价[A01,A10,B01,B02]。它们说明评价器本身可以构成研究工作，但训练一个新评分头或增加几个维度不足以自然建立差异。阅读重点应从“谁分数最高”转到标签定义、切分与适用范围。'),
('从整体结果走向局部证据','ImageDoctor、SpatialReward、BPM及EditInspector将定位、语义核验和细粒度错误纳入评价[A02,A03,B11,B12]；PVC-Judge和AugCLIP强调保持与修改[B03,B13]。因此“目标改好与背景改坏分开评价”是基础要求。可进一步检验的是每条诊断的证据是否充分，以及诊断对有效修正和无关变化的响应。'),
('从解释文本走向验证与检查','ReasonEdit用解释质量监督训练评价，EVR引入假设与证据验证，RewardHarness选择并演化工具技能，CheckDistill把检查项知识蒸馏给高效评分器[A04,A05,A07,A09]。这些方向已覆盖解释、验证和检查清单，不能统称为只会静态打分。我们的候选差异必须落在状态依赖的逐次检查价值和可检验诊断响应，而非流程名称。'),
('从评价走向纠错效果','R³-Refiner、SQUARE-Bench和UMM-Reflection分别研究反馈、固定编辑器闭环和统一模型的反思修正[A06,B05,B06]。正确诊断不保证工具能修好，修复后总分上升也不能证明反馈正确。应将事实正确性、旧错误减少、新损伤和检查成本分开测量。')
]
for title,body in review:h(d,title);p(d,body);md.extend(['## '+title,body])
page(d,'初步综述与可验证假设')
for title,body in [
('反事实和假设检验也有前作','EditJudgeBias通过质量保持审核研究无关因素偏差[A08]；MDC已用属性编辑进行模型假设检验[B16]。后者研究视觉识别模型诊断，任务与我们不同，但方法思想不能忽略。我们的主张应是针对编辑要求、诊断响应和有限成本证据获取的具体学习机制，不是首次使用反事实。'),
('目前值得验证的缺口','H1：等人工成本下，修正与干扰对照监督能否改善真实编辑诊断，而不只改善合成样例。H2：同骨干、同工具、同预算下，基于当前疑点的检查选择能否优于固定清单和直接工具路由。H3：将该诊断输入队友的固定决策方案后，是否改善最终编辑结果或节省成本。这些是待验证假设，不是已证明的文献空白。'),
('建议综述题目与结构','建议题目：图像编辑评价中的局部证据 诊断可靠性与主动核验。正文按五部分写：任务与评价维度；人工数据及专用评价模型；定位解释与证据验证；主动检查及高效蒸馏；反馈纠错与可靠性审计。最后提出三项假设。初稿目标约3000–5000字，以问题组织，避免按30篇逐篇堆摘要。'),
('读完后必须留下的证据','每篇记录输入、输出、训练对象、监督来源、证据获取方式、成本、数据切分、最强基线、已证实能力、未回答问题，以及支持判断的页码或章节。把“作者说”“实验支持”“我推测”分开。没有主动查证过的能力不要写成“该论文没有”。')]:h(d,title);p(d,body);md.extend(['## '+title,body])

page(d,'十二周阅读与小实验安排')
schedule=[['1–2','A01 A02；B01 B02 B11 B12','画评价器输入输出，跑5个现成样例'],['3–4','A03 A04；B03 B04 B09 B10','填区域与解释比较表，试标10组'],['5–6','A05 A07；B13 B14 B16 B18','画工具路由与蒸馏流程，核查创新重合'],['7–8','A08 A09；B15 B17 B19 B20','设计有效干预与噪声对照，完成20组试标'],['9–10','A06 A10；B05 B06 B07 B08','区分诊断真值与修复收益，写数据计划'],['11–12','回读关键方法与附录','完成30篇矩阵和综述初稿，与导师讨论两个假设']]
table(d,['周次','阅读','同时完成的产物'],schedule,[2,6,8])
p(d,'前12周每周阅读约4–5小时，另留1–2小时小实验与讨论。不要读完30篇才动手。余下年度按问题回读，每月更新一次检索，新增文章替换低相关条目，不无限扩充清单。该安排不自动建立后台监控。')
h(d,'轻量阅读卡模板')
p(d,'完整题名与版本：____；一句话问题：____；方法输入输出：____；训练与监督：____；证据与检查：____；主要实验及支持页码：____；可复用资源：____；与我重合：____；可验证差异：____；本轮仍不理解：____。')
p(d,'只挑一个专用评价基线和一个主动/工具基线深入运行，其余按相关程度读方法和实验。代码公开不等于权重可用，摘要称将发布不等于已发布；使用数据前核实许可、访问条件和原图来源。')
md.extend(['## 12周阅读路径']+[' / '.join(r) for r in schedule])
lit=OUT/'htz_必读10篇与了解20篇_文献指南和综述提纲_20261028.docx';d.save(lit)
(OUT/'htz_文献30篇阅读清单_20261028.md').write_text('\n\n'.join(md),encoding='utf-8')
(BASE/'literature30.json').write_text(json.dumps({'A':A,'B':B},ensure_ascii=False,indent=2),encoding='utf-8')

d=setup('黄天志一年期精简研究计划')
p(d,'保留可检验诊断与主动检查两项核心\n新版本规划 · 原申请表和前版方案均保留 · 2026年10月28日')
h(d,'结论与适用条件')
p(d,'建议按一年日历周期，安排46个有效工作周和6周考试、休息、设备等待缓冲。本人平均每个有效周约6–7小时，基准总量296小时；按全年52周摊约5.7小时。不是每周都均匀工作，训练与定稿周可能需8–10小时，考试周可0–3小时。')
p(d,'总量目标由前版300–400小时调整为250–310小时，其中独立课题220–270小时，团队额外30–40小时。296小时基准比前版360小时少约18%。这是通过减少部署与扩展任务得到的条件性预算，不保证复杂研究都能压缩；若团队服务量不减或基础学习超过预期，应回到旧范围。')
table(d,['工作类别','基准小时','压缩方式'],[['文献 学习与综述','60','10篇精读20篇定向读，只复现少数关键方法'],['环境及基础工具','24','一个骨干，最多3类检查，统一缓存'],['数据与标注','40','先20组测时，复用许可数据，短结构化标签'],['可检验诊断训练','44','小头优先，一条主训练路线'],['主动检查学习','36','有限候选区域和工具，监督排序优先'],['独立验证与统计','36','预设必要消融，留出真实测试'],['报告 代码与复现','20','边实验边记录，复用统计脚本'],['个人合计','260','包含文献与标注，不另重复增加'],['团队额外职责','36','接口 校准配额 组织与统计，需共同确认'],['总计','296','不含其他标注者、队友和GPU被动等待']],[5,2,9])
p(d,'40小时数据预算含构造审核和本人标注，是偏紧的目标，不能与旧版“本人标注40–60小时”同时承诺。须通过20组试标确认；若不够，优先减少扩展类别或将数据阶段提高至50–60小时，总预算相应增加。不能省掉独立真值审核。')

page(d,'一年时间线与每周安排')
table(d,['阶段','有效周','本人研究','团队','产物'],[['学习与先导','1–12','72','6','30篇矩阵 综述初稿 20组试标'],['数据与工具','13–20','48','6','拆分冻结 数据集v1 缓存工具'],['诊断学习','21–28','44','6','模型v1 真实诊断与响应结果'],['主动核验','29–36','40','6','下一步区域与工具选择 消融'],['独立测试','37–42','36','6','泛化 故障 成本与置信区间'],['总结与接口','43–46','20','6','研究报告 复现包 团队接入'],['合计','46周','260','36','296小时']],[3.3,2,2.1,1.5,7.1])
p(d,'有效周不等于连续日历周。示例日历：2026年11月—2027年1月学习和试跑；2—3月数据；4—5月诊断训练；6—7月主动核验；8—9月验证；10月总结。6个缓冲周分散在考试、节假日和设备等待期间。实际起止随学校申报和导师安排调整，不代表正式截止时间。')
h(d,'每周如何安排')
p(d,'前期每周约4–5小时阅读、1–2小时小实验或讨论。中后期可用两个2小时整块做实现和实验，1小时读相关论文，约1小时整理数据与记录，再以隔周1小时处理团队事务。阅读时间从总预算内分配，不在每天额外叠加。')
p(d,'每四个有效周检查一次实际工时、标注速度与完成项。若连续两个月超预算，先删扩展，不删真值与基线。若每周最多5小时，46周只有230小时，低于本方案250小时下限；应延长周期、获得协助，或明确缩小成果范围。')
h(d,'学习和实验交错')
p(d,'第2周即跑少量样例，第4周形成失败分类，第8周完成对照有效性审核，第12周向导师提交初稿和预实验。不要把前半年全部用于读论文，也不要读完所有文章才建立数据管线。')

page(d,'压缩边界与团队承诺')
table(d,['必须保留','推迟或复用'],[['可检验诊断与下一步检查学习','不额外开发完整修复器'],['固定清单与同预算强基线','不部署所有开源评价器'],['真实性与干预有效性审核','不要求逐图长推理链或像素精标'],['真实留出测试和泛化','第二编辑器先验证，不完整复制所有生成实验'],['数据隔离 成本 不确定性','暂不做多轮历史和复杂变化依赖图'],['独立个人报告与团队接口','大规模反馈RL、在线反事实生成列为扩展']],[8,8])
p(d,'原有28个idea仍保留在全景文件中；本年度核心是诊断响应学习与主动检查，其他方向作为场景、支撑能力或未来扩展。减少工程铺设不等于把研究缩成打分器。')
h(d,'团队工时降低的前提')
p(d,'团队校准开发标注三人轮换，最终组外评分由团队共同落实；你负责协议与统计但不包办所有评分。原计划34–164合计人时的人工评价不能隐藏在个人36小时团队预算里。若团队仍要求你承担原40–60小时，则新总量应约260–330小时或更高，而不是默默压缩独立研究。')
h(d,'不能预先承诺的部分')
p(d,'训练数据达到何种规模才有效、跨编辑器是否泛化、API费用和实验室GPU支持仍待实测。预算按冻结骨干和监督学习估算；如必须大模型微调或多轮RL，先重新核算。论文接收与“首次”不在时间表承诺中。')
p(d,'建议导师讨论时确认：一年研究期限、每周6–7小时是否可接受、是否支持固定骨干路线、标注助手和组外评审安排、何时进行中期审查。年度版本与前版工时不同，应采用本版预算而非把两版相加。')
annual=OUT/'htz_一年期精简研究计划与工时预算_20261028.docx';d.save(annual)
print(json.dumps({'literature_count':len(A)+len(B),'core':len(A),'awareness':len(B),'outputs':[str(lit),str(annual)]},ensure_ascii=False))
